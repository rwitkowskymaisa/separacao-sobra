"""Gera dados/plano.json a partir da planilha do plano de separação (Cruzamento_Sobra_EX_x_CBP_CBTBE_vN.xlsx).

Uso:
    pip install openpyxl
    python scripts/gerar_plano_json.py caminho/da/planilha.xlsx v5-2026-11-03

A planilha precisa ter as abas Caixas, Itens, Pedido x Sobra e Complemento CBP, como a versão 4,
e estar salva com os valores calculados (abra e salve no Excel antes, se tiver editado).
Depois: commit do dados/plano.json no GitHub e, no sistema, "Carregar plano do repositório".
"""
import json, sys
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

def linhas(ws):
    it = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c is not None else '' for c in next(it)]
    for r in it:
        if r and any(v is not None for v in r):
            yield dict(zip(cab, r))

def txt(v):
    return '' if v is None else str(v).strip()

def main(xlsx, versao):
    wb = load_workbook(xlsx, data_only=True)
    caixas = {}
    for r in linhas(wb['Caixas']):
        cid = txt(r['Nº Caixa'])
        if not cid: continue
        caixas[cid] = {'sit': txt(r['Situação']), 'notas': [int(n) for n in txt(r['Notas']).replace(' ', '').split(',') if n],
                       'comp': int(r.get('Complemento CBP (un)') or 0)}
    ordem = [txt(r['Nº Caixa']) for r in sorted(linhas(wb['Complemento CBP']), key=lambda r: r['Ordem'] or 0) if r.get('Nº Caixa')]
    itens = []
    for r in linhas(wb['Itens']):
        if not r.get('ISBN'): continue
        vals = [r['Qtd p/ CBP'], r['Qtd p/ CBTBE'], r['Qtd devolução']]
        if any(v is None for v in vals):
            sys.exit('A aba Itens está sem valores calculados. Abra a planilha no Excel, salve e rode de novo.')
        itens.append({'isbn': txt(r['ISBN']), 'desc': txt(r['Descrição']), 'caixa': txt(r['Nº Caixa']), 'nota': int(r['Nota']),
                      'est': int(r['Estoque']), 'C': int(vals[0]), 'E': int(vals[1]), 'D': int(vals[2]), 'K': 0})
    pedidos = []
    for r in linhas(wb['Pedido x Sobra']):
        if not r.get('ISBN'): continue
        pedidos.append({'isbn': txt(r['ISBN']), 'tit': txt(r['Título']), 'autor': txt(r.get('Autor')) or None,
                        'ev': txt(r['Evento']), 'qtd': int(r['Qtd pedido'])})
    ped_cbp = {p['isbn']: p['qtd'] for p in pedidos if p['ev'] == 'CBP'}
    # complemento por linha: excesso do CBP sobre o pedido, atribuído às caixas convertidas (da última para a primeira)
    tot = defaultdict(int)
    for i in itens: tot[i['isbn']] += i['C']
    for isbn, t in tot.items():
        resto = t - ped_cbp.get(isbn, 0)
        for cx in reversed(ordem):
            if resto <= 0: break
            for i in itens:
                if i['isbn'] == isbn and i['caixa'] == cx and resto > 0:
                    i['K'] = min(resto, i['C']); resto -= i['K']
        if resto > 0:
            print(f'Aviso: complemento de {resto} un do ISBN {isbn} não ficou em caixa convertida.')
    desc = {}
    for i in itens: desc.setdefault(i['isbn'], i['desc'])
    tit = {p['isbn']: p['tit'] for p in pedidos}
    isbns = sorted(set(desc) | set(tit)); ix = {k: n for n, k in enumerate(isbns)}
    T = [[k, desc.get(k, ''), '' if tit.get(k, '').upper() == desc.get(k, '').upper() else tit.get(k, '')] for k in isbns]
    B = []
    for cid in sorted(caixas, key=lambda c: (c == 'S/Caixa', c)):
        c = caixas[cid]
        rows = [[ix[i['isbn']], i['nota'], i['est'], i['C'], i['K'], i['E'], i['D']] for i in itens if i['caixa'] == cid]
        B.append([cid, c['sit'], 1 if c['comp'] > 0 else 0, c['notas'], rows])
    Pd = [[ix[p['isbn']], p['ev'], p['qtd'], p['autor']] for p in pedidos]
    out = {'versao': versao, 't': T, 'b': B, 'p': Pd}
    dest = Path(__file__).resolve().parent.parent / 'dados' / 'plano.json'
    dest.write_text(json.dumps(out, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'{dest}: {len(B)} caixas, {len(itens)} itens, {len(Pd)} linhas de pedido, complemento {sum(i["K"] for i in itens)} un')

if __name__ == '__main__':
    if len(sys.argv) != 3: sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
